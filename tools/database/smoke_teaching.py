"""End-to-end teaching API checks against a local running server.

Uses local initialization credentials; cleans only its explicitly created test
records in finally. MYSQL_PWD is used exclusively for that test-data cleanup.
"""
from io import BytesIO
import json
import os
from pathlib import Path
import sys
from urllib.request import Request, urlopen
from urllib.parse import urlencode
import uuid

import pymysql
from openpyxl import Workbook, load_workbook

BASE = os.environ.get('TEACHING_TEST_URL', 'http://127.0.0.1:8080/springboote51e2')
ROOT = Path(__file__).resolve().parents[2]


def call(path, token=None, method='GET', body=None, form=False, raw=False):
    headers = {'Token': token} if token else {}
    if body is not None:
        headers['Content-Type'] = 'application/x-www-form-urlencoded' if form else 'application/json'
        body = (urlencode(body) if form else json.dumps(body)).encode('utf-8')
    with urlopen(Request(BASE + path, body, headers, method=method), timeout=40) as response:
        content = response.read()
    return content if raw else json.loads(content)


def upload(path, token, content, filename='test.xlsx'):
    boundary = '----TeachingTest' + uuid.uuid4().hex
    data = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{filename}"\r\n'
            'Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n').encode() + content + f'\r\n--{boundary}--\r\n'.encode()
    with urlopen(Request(BASE+path, data, {'Token':token, 'Content-Type':'multipart/form-data; boundary='+boundary}, method='POST'), timeout=60) as response:
        return json.load(response)


def ok(result):
    assert result.get('code') == 0, result
    return result.get('data')


def main():
    accounts = json.loads((ROOT/'database/generated/teaching-accounts.local.json').read_text(encoding='utf-8'))
    tokens = []
    def login(account):
        path = '/users/login' if account['role']=='管理员' else '/jiaoshi/login'
        result = call(path, method='POST', body=dict(username=account['username'],password=account['password']), form=True)
        assert result.get('code') == 0, 'Login failed'
        tokens.append(result['token'])
        return result['token']
    admin = login(accounts[0])
    lookup = ok(call('/teaching/lookups', admin))
    current = next(t for t in lookup['terms'] if t['status']=='OPEN')
    teacher = lookup['teachers'][0]
    account = next(a for a in accounts if a['username']==teacher['gonghao'])
    own_token = login(account)
    other_account = next(a for a in accounts if a['role']=='教师' and a['username']!=teacher['gonghao'])
    other_token = login(other_account)
    checks, task_ids, lab_ids, teacher_ids = [], [], [], []
    marker = 'SMOKE-' + uuid.uuid4().hex[:12]
    try:
        assert call('/teaching/dashboard')['code']==401
        for old in ['/users/resetPass?username=admin', '/jiaoshi/page', '/shiyankecheng/list']:
            assert call(old,admin)['code'] != 0, old
        for path in ['/teaching/../users/list', '/teaching/%2e%2e/users/list']:
            assert call(path, other_token)['code'] == 404, 'Resolved legacy route must remain disabled'
        dashboard=ok(call('/teaching/dashboard',admin))
        assert dashboard['counts']['tasks'] >= 255
        checks.append('管理员登录、真实概览、未登录拒绝与旧入口关闭')

        tasks=ok(call('/teaching/tasks?limit=200',admin))
        assert tasks['total'] >= 255
        historical=next(t for t in tasks['list'] if t['status']=='ARCHIVED')
        own=ok(call('/teaching/tasks?limit=200',own_token))
        for task in own['list']:
            detail=ok(call('/teaching/tasks/'+str(task['id']),own_token))
            assert teacher['id'] in detail['teacher_ids']
        visible_teachers=ok(call('/teaching/teachers',other_token))['list']
        assert len(visible_teachers)==1 and visible_teachers[0]['gonghao']==other_account['username']
        assert 'mima' not in visible_teachers[0]
        assert call('/teaching/teachers',other_token,'POST',dict(gonghao=marker,jiaoshixingming='越权'))['code'] != 0
        checks.append('教师课程范围及管理员接口隔离')

        newtask=dict(termId=current['id'],courseId=lookup['courses'][0]['id'],
                     teacherIds=[teacher['id']],labId=lookup['labs'][0]['id'],
                     classComposition=marker,enrollmentCount=47,plannedLabHours=4,majorComposition='接口验证')
        created=ok(call('/teaching/tasks',admin,'POST',newtask)); task_ids.append(created['id'])
        task_id=created['id']
        assert call('/teaching/tasks/'+str(task_id),other_token)['code'] != 0
        assert ok(call('/teaching/projects?taskId='+str(task_id),own_token))['editable'] is True
        project=dict(task_id=task_id,school_code='11059',name=marker,category_code='3',type_code='2',
                     discipline_code='0809',requirement_code='1',participant_type_code='3',group_size=1,hours=2,sort_order=1)
        added=ok(call('/teaching/projects',own_token,'POST',project))
        project_id=added['id']
        assert call('/teaching/projects/'+str(project_id),other_token,'PUT',project)['code'] != 0
        project['name']=marker+'修改'
        ok(call('/teaching/projects/'+str(project_id),own_token,'PUT',project))
        assert ok(call('/teaching/projects?taskId='+str(task_id),own_token))['list'][0]['name']==project['name']
        historical_project=dict(project,task_id=historical['id'])
        assert call('/teaching/projects',admin,'POST',historical_project)['code'] != 0
        checks.append('当前任务创建、教师项目增改、越权拒绝、历史学期只读')

        second=ok(call('/teaching/tasks',admin,'POST',dict(newtask,classComposition=marker+'复制')))
        task_ids.append(second['id'])
        ok(call('/teaching/projects/copy',own_token,'POST',dict(sourceTaskId=task_id,targetTaskId=second['id'])))
        copied=ok(call('/teaching/projects?taskId='+str(second['id']),own_token))
        assert copied['total']==1
        assert call('/teaching/projects/'+str(project_id),own_token,'DELETE')['code'] != 0
        ok(call('/teaching/projects/'+str(copied['list'][0]['id']),own_token,'DELETE'))
        ok(call('/teaching/projects/'+str(project_id),own_token,'DELETE'))
        assert ok(call('/teaching/projects?taskId='+str(task_id),own_token))['total']==0
        checks.append('项目复制与删除')

        templates={}
        for name in ['projects','teachers','labs','timetable']:
            binary=call('/teaching/templates/'+name,admin,raw=True)
            book=load_workbook(BytesIO(binary)); templates[name]=[c.value for c in book.worksheets[0][1]]
            assert templates[name]
        book=Workbook();sheet=book.active;sheet.append(templates['projects'])
        course=next(c for c in lookup['courses'] if c['id']==newtask['courseId'])
        sheet.append(['11059','',marker+'导入','3','2','0809','1','3','',1,2,course['course_code'],course['course_name'],4,teacher['jiaoshixingming']])
        buf=BytesIO();book.save(buf)
        ok(upload('/teaching/imports/projects?taskId='+str(task_id),own_token,buf.getvalue()))
        assert ok(call('/teaching/projects?taskId='+str(task_id),own_token))['total']==1
        checks.append('四类Excel模板可下载、项目Excel真实导入')

        year_id=next(t['academic_year_id'] for t in lookup['terms'] if t['name'].startswith('2025-2026'))
        reports=ok(call('/teaching/reports?yearId='+str(year_id),admin))
        assert len(reports['labs'])==14
        for report_type in ['labs','projects']:
            binary=call('/teaching/reports/export?yearId='+str(year_id)+'&type='+report_type,admin,raw=True)
            assert load_workbook(BytesIO(binary)).worksheets[0].max_row>=1
        status=ok(call('/teaching/ai/status',admin))
        if not status['configured']:
            assert call('/teaching/ai/query',admin,'POST',{'question':'统计实验室人时数'})['code']!=0
        assert call('/teaching/ai/query',other_token,'POST',{'question':'显示所有教师'})['code']!=0
        checks.append('14间实验室报表、两类Excel导出、AI未配置/权限状态真实反馈')

        # Ordinary maintenance can create/update independent records without altering real teacher data.
        lab=ok(call('/teaching/labs',admin,'POST',dict(shiyanshibianhao=marker,shiyanshimingcheng='验证实验室',shiyanshiweizhi='验证地点',equipment_count=12)))
        lab_ids.append(lab['id'])
        ok(call('/teaching/labs/'+str(lab['id']),admin,'PUT',dict(shiyanshibianhao=marker,shiyanshimingcheng='验证实验室改',shiyanshiweizhi='验证地点',equipment_count=13)))
        teacher_record=ok(call('/teaching/teachers',admin,'POST',dict(gonghao=marker,jiaoshixingming='验证教师',xueyuan='测试')))
        teacher_ids.append(teacher_record['id'])
        reset=ok(call('/teaching/teachers/'+str(teacher_record['id'])+'/reset-password',admin,'POST',{}))
        assert len(reset['password'])>=8
        reset_token=login(dict(role='教师', username=marker, password=reset['password']))
        ok(call('/jiaoshi/logout',reset_token,'POST'))
        assert call('/teaching/dashboard',reset_token)['code']==401
        checks.append('实验室维护、教师创建与密码重置')
        checks.append('特殊路径仍禁用旧接口、密码重置后可登录、退出使Token失效')
        result=dict(status='PASS',checks=checks,temporary_records_removed=True)
    finally:
        connection=pymysql.connect(host='127.0.0.1',user='root',password=os.environ['MYSQL_PWD'],database='t132',charset='utf8mb4')
        try:
            with connection.cursor() as cursor:
                for task_id in reversed(task_ids):
                    cursor.execute('UPDATE experiment_project SET copied_from_id=NULL WHERE task_id=%s',(task_id,))
                    cursor.execute('DELETE FROM experiment_project WHERE task_id=%s',(task_id,))
                    cursor.execute('DELETE FROM schedule_detail WHERE task_id=%s',(task_id,))
                    cursor.execute('DELETE FROM teaching_task_teacher WHERE task_id=%s',(task_id,))
                    cursor.execute('DELETE FROM teaching_task WHERE id=%s',(task_id,))
                for lab_id in lab_ids: cursor.execute('DELETE FROM shiyanshixinxi WHERE id=%s',(lab_id,))
                for teacher_id in teacher_ids:
                    cursor.execute("DELETE FROM token WHERE tablename='jiaoshi' AND userid=%s",(teacher_id,))
                    cursor.execute('DELETE FROM jiaoshi WHERE id=%s',(teacher_id,))
                for token in tokens: cursor.execute('DELETE FROM token WHERE token=%s',(token,))
            connection.commit()
        finally: connection.close()
    (ROOT/'docs/database/reports/teaching-api-smoke.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__': main()

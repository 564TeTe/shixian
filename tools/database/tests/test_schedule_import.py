import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
from datetime import time
from openpyxl import Workbook

MODULE = Path(__file__).resolve().parents[1] / 'schedule_import.py'
if MODULE.exists():
    spec = importlib.util.spec_from_file_location('schedule_import', MODULE)
    subject = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = subject
    spec.loader.exec_module(subject)
else:
    subject = None


class ScheduleTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(subject, '课表解析模块尚未实现')

    def test_multiple_labs_person_hours(self):
        rows = subject.parse_schedule('36-506;36-508',
            '星期二第5-8节{6-7周};星期二第5-8节{8周}')
        totals = {}
        for row in rows:
            totals[row['lab_code']] = totals.get(row['lab_code'], 0) + row['hours'] * 47
        self.assertEqual(totals, {'36-506': 376, '36-508': 188})

    def test_odd_weeks_and_discontinuous_periods(self):
        rows = subject.parse_schedule('36-401', '星期六第1-2,5-6节{7-11周(单),14周}')
        self.assertEqual(sorted({r['week'] for r in rows}), [7, 9, 11, 14])
        self.assertEqual(len(rows), 8)
        self.assertEqual(sum(r['hours'] for r in rows), 16)

    def test_even_weeks(self):
        rows = subject.parse_schedule('36-401', '星期一第3-4节{10-15周(双)}')
        self.assertEqual([r['week'] for r in rows], [10, 12, 14])

    def test_invalid_or_mismatched_schedule_is_rejected(self):
        cases = [
            ('36-401;36-403', '星期一第1-2节{1周}'),
            ('36-401', '星期一第1-2节{9-3周}'),
            ('36-401', '星期一第0-2节{1周}'),
            ('36-401', '星期一第1-2节{1-2周(单),未知}'),
            ('36-401', '星期一第1-2节{0周}'),
            ('36-401', '星期一第1-2,2-3节{1周}'),
        ]
        for locations, times in cases:
            with self.subTest(times=times), self.assertRaises(ValueError):
                subject.parse_schedule(locations, times)

    def test_nonnegative_integer_validation(self):
        self.assertEqual(subject.number('35', integer=True), 35)
        self.assertEqual(subject.number('0', integer=True), 0)
        for value in ['-1', '3.5', 'NaN', 'Infinity', '', None]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                subject.number(value, integer=True)

    def test_sql_text_handles_quotes_and_backslashes(self):
        encoded = subject.sql_text("a'\\b\n中文")
        self.assertEqual(encoded, "CONVERT(0x61275c620ae4b8ade69687 USING utf8mb4)")

    def analyze_rows(self, rows):
        workbook = Workbook()
        workbook.active.append(subject.HEADERS)
        for row in rows:
            workbook.active.append(row)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'sample.xlsx'
            workbook.save(path)
            return subject.analyze(path)

    def sample(self):
        return ['2025-2026', '1', '学院', '00123', '实验课', '2', '班1', '40',
                '老师甲', '2', '1-2周', '4', '36-401', None, '35', '1', '2',
                '1-2周', '2025-12-16 ', '实验(2.0)', '星期二第1-2节{1-2周}']

    def test_real_file_date_with_trailing_space_keeps_raw_value(self):
        records, report = self.analyze_rows([self.sample()])
        self.assertNotIn('INVALID', report['issue_counts'])
        self.assertEqual(records[0]['raw']['课程结束时间'], '2025-12-16 ')
        self.assertEqual(records[0]['raw']['课程号'], '00123')
        self.assertEqual(report['expanded_schedule_rows'], 2)

    def test_ambiguous_task_rows_are_preserved(self):
        first, second = self.sample(), self.sample()
        second[-1] = '星期三第1-2节{1-2周}'
        records, report = self.analyze_rows([first, second])
        self.assertEqual(report['row_count'], 2)
        self.assertEqual(report['issue_counts']['TASK_IDENTITY_AMBIGUOUS'], 2)
        self.assertEqual([r['source_row'] for r in records], [2, 3])

    def test_invalid_numeric_row_is_retained_for_review(self):
        row = self.sample()
        row[14] = '-2'
        records, report = self.analyze_rows([row])
        self.assertEqual(report['row_count'], 1)
        self.assertEqual(report['issue_counts']['INVALID'], 1)
        self.assertIsNone(records[0]['schedule'])

    def test_overlapping_time_across_labs_is_rejected(self):
        with self.assertRaises(ValueError):
            subject.parse_schedule('36-401;36-403',
                '星期二第1-2节{1周};星期二第2-3节{1周}')

    def test_data_sheet_with_missing_header_cannot_be_silently_skipped(self):
        workbook = Workbook()
        workbook.active.append(subject.HEADERS)
        workbook.active.append(self.sample())
        malformed = workbook.create_sheet('漏表头')
        malformed.append([None] * 21)
        malformed.append(self.sample())
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'sample.xlsx'
            workbook.save(path)
            with self.assertRaises(ValueError):
                subject.analyze(path)

    def test_time_only_cell_can_still_be_staged_as_invalid(self):
        row = self.sample()
        row[18] = time(8, 30)
        records, report = self.analyze_rows([row])
        self.assertEqual(report['issue_counts']['INVALID'], 1)
        self.assertIsInstance(records[0]['raw']['课程结束时间'], str)
        sql = subject.staging_sql(records, report)
        self.assertIn('INSERT INTO teaching_import_row', sql)


if __name__ == '__main__':
    unittest.main()

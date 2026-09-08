"""Validate timetable workbooks and generate SQL for review-only staging."""
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook

PARSER_VERSION = '1.0.0'
HEADERS = ['学年', '学期', '开课学院', '课程号', '课程名称', '学分',
           '教学班组成', '教学班人数', '教师名称', '周学时', '起始结束周',
           '课程实验总学时', '教学地点', '专业组成', '选课人数', '起始周',
           '结束周', '排课起始结束周', '课程结束时间', '课程周学时', '上课时间']


def number(value, integer=False):
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f'不是数值: {value!r}') from None
    if not result.is_finite() or result < 0:
        raise ValueError(f'必须是有限非负数: {value!r}')
    if integer:
        if result != result.to_integral_value() or result > 2147483647:
            raise ValueError(f'必须是有效非负整数: {value!r}')
        return int(result)
    if result > Decimal('999999.99') or result != result.quantize(Decimal('.01')):
        raise ValueError(f'数值超出范围或超过两位小数: {value!r}')
    return result


def _range(token, maximum):
    if not re.fullmatch(r'\d+(?:-\d+)?', token):
        raise ValueError(f'非法范围: {token}')
    edges = [int(n) for n in token.split('-')]
    start, end = edges[0], edges[-1]
    if not 1 <= start <= end <= maximum:
        raise ValueError(f'范围越界或逆序: {token}')
    return set(range(start, end + 1))


def parse_schedule(locations, times):
    places, parts = str(locations).split(';'), str(times).split(';')
    if len(places) != len(parts):
        raise ValueError('教学地点与上课时间的片段数不一致')
    result, occupied = [], set()
    for index, (place, part) in enumerate(zip(places, parts), 1):
        place = place.strip()
        match = re.fullmatch(r'星期([一二三四五六日天])第([\d,-]+)节\{(.+)\}', part.strip())
        if not place or place in ('None', '无') or not match:
            raise ValueError(f'无法识别地点或时间片段 {index}')
        day, periods_text, weeks_text = match.groups()
        weekday = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '日': 7, '天': 7}[day]
        periods, weeks = set(), set()
        for token in periods_text.split(','):
            segment = _range(token, 24)
            if periods & segment:
                raise ValueError('同一片段节次重叠')
            periods.update(segment)
        for token in weeks_text.split(','):
            wm = re.fullmatch(r'([\d-]+)周(?:\(([单双])\))?', token)
            if not wm:
                raise ValueError(f'无法识别周次: {token}')
            selected = _range(wm[1], 53)
            if wm[2]:
                selected = {w for w in selected if w % 2 == (1 if wm[2] == '单' else 0)}
            if weeks & selected:
                raise ValueError('同一片段周次重复')
            weeks.update(selected)
        if not weeks:
            raise ValueError('周次筛选后为空')
        intervals = []
        for period in sorted(periods):
            if intervals and intervals[-1][1] + 1 == period:
                intervals[-1][1] = period
            else:
                intervals.append([period, period])
        for week in sorted(weeks):
            for period in periods:
                key = (week, weekday, period)
                if key in occupied:
                    raise ValueError('同一任务的排课片段在时间上重叠，需人工核对')
                occupied.add(key)
            for first, last in intervals:
                result.append(dict(lab_code=place, week=week, weekday=weekday,
                                   period_start=first, period_end=last,
                                   hours=last-first+1, source_segment=index))
    return result


def sql_text(value):
    encoded = str(value).encode('utf-8').hex()
    return f'CONVERT(0x{encoded} USING utf8mb4)' if encoded else "''"


def _json_value(value):
    if isinstance(value, datetime):
        return value.isoformat(sep=' ')
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value) if isinstance(value, timedelta) else value


def analyze(path):
    path = Path(path)
    workbook = load_workbook(path, data_only=False, read_only=True)
    records, identity_groups = [], defaultdict(list)
    names, course_codes, labs = set(), set(), set()
    course_names = defaultdict(set)
    try:
        for sheet in workbook:
            values = sheet.iter_rows(values_only=True)
            header = [str(v).strip() if v is not None else '' for v in next(values, [])]
            if not any(header):
                if any(any(v is not None for v in row) for row in values):
                    raise ValueError(f'{sheet.title}: 首行缺少表头，但后续存在数据，不能忽略')
                continue
            if header != HEADERS:
                raise ValueError(f'{sheet.title}: 表头必须与21列课表模板一致')
            for row_number, row in enumerate(values, 2):
                if not any(v is not None for v in row):
                    continue
                raw = dict(zip(HEADERS, map(_json_value, row)))
                issues, schedule = [], None
                issues.append('TEACHER_ID_MAPPING_REQUIRED')
                try:
                    for key in HEADERS:
                        if key != '专业组成' and (raw[key] is None or not str(raw[key]).strip()):
                            raise ValueError(f'{key}不能为空')
                        if isinstance(raw[key], str) and raw[key].startswith('='):
                            raise ValueError(f'{key}包含公式，需核对并提供原始值')
                    year_match = re.fullmatch(r'(\d{4})-(\d{4})', str(raw['学年']))
                    if not year_match or int(year_match[2]) != int(year_match[1]) + 1:
                        raise ValueError('学年格式不正确')
                    if number(raw['学期'], integer=True) not in (1, 2):
                        raise ValueError('学期必须为1或2')
                    for key in ['教学班人数', '选课人数', '起始周', '结束周']:
                        number(raw[key], integer=True)
                    for key in ['学分', '周学时', '课程实验总学时']:
                        number(raw[key])
                    if not 1 <= number(raw['起始周'], integer=True) <= number(raw['结束周'], integer=True) <= 53:
                        raise ValueError('起始周或结束周不合法')
                    if number(raw['课程实验总学时']) == 0:
                        raise ValueError('实验总学时必须大于0')
                    datetime.fromisoformat(str(raw['课程结束时间']).strip())
                    schedule = parse_schedule(raw['教学地点'], raw['上课时间'])
                except ValueError as exc:
                    issues.append('INVALID: ' + str(exc))
                if schedule is not None:
                    if number(raw['课程实验总学时']) != sum(s['hours'] for s in schedule):
                        issues.append('HOURS_MISMATCH')
                    if number(raw['教学班人数'], integer=True) != number(raw['选课人数'], integer=True):
                        issues.append('HEADCOUNT_DIFFERS')
                    labs.update(s['lab_code'] for s in schedule)
                teachers = [n.strip() for n in re.split('[,，;；、]', str(raw['教师名称'] or '')) if n.strip()]
                names.update(teachers)
                if len(teachers) > 1:
                    issues.append('CO_TEACHING')
                if raw['专业组成'] is None or not str(raw['专业组成']).strip():
                    issues.append('MAJOR_MISSING')
                code = str(raw['课程号'] or '').strip()
                course_codes.add(code)
                course_names[code].add(str(raw['课程名称']))
                record = dict(sheet=sheet.title, source_row=row_number, raw=raw,
                              schedule=schedule, issues=issues)
                records.append(record)
                key = tuple(str(raw[k]) for k in ['学年', '学期', '课程号', '教学班组成']) + (tuple(sorted(teachers)),)
                identity_groups[key].append(record)
    finally:
        workbook.close()
    if not records:
        raise ValueError('没有可暂存的课表数据')
    duplicates = []
    for group in identity_groups.values():
        if len(group) > 1:
            duplicates.append([{'sheet': r['sheet'], 'row': r['source_row']} for r in group])
            for record in group:
                record['issues'].append('TASK_IDENTITY_AMBIGUOUS')
    for record in records:
        if len(course_names[str(record['raw']['课程号'] or '').strip()]) > 1:
            record['issues'].append('COURSE_NAME_CONFLICT')
    report = dict(file=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                  parser_version=PARSER_VERSION, row_count=len(records),
                  course_code_count=len(course_codes), teacher_name_count=len(names),
                  lab_code_count=len(labs),
                  raw_schedule_segments=sum(len(str(r['raw']['上课时间']).split(';')) for r in records),
                  expanded_schedule_rows=sum(len(r['schedule'] or []) for r in records),
                  terms=dict(Counter(str(r['raw']['学期']) for r in records)),
                  issue_counts=dict(Counter(i.split(':')[0] for r in records for i in r['issues'])),
                  ambiguous_task_groups=duplicates,
                  hours_mismatches=[dict(sheet=r['sheet'], row=r['source_row'],
                                        planned=r['raw']['课程实验总学时'],
                                        scheduled=sum(s['hours'] for s in r['schedule']))
                                    for r in records if 'HOURS_MISMATCH' in r['issues']],
                  row_issues=[dict(sheet=r['sheet'], row=r['source_row'], issues=r['issues']) for r in records])
    return records, report


def staging_sql(records, report):
    text = ['-- Generated review-only staging. Contains source workbook data.',
            'SET NAMES utf8mb4;', 'START TRANSACTION;',
            'INSERT INTO teaching_import_batch (file_sha256,file_name,parser_version,row_count) VALUES (' +
            ','.join([sql_text(report['sha256']), sql_text(report['file']),
                      sql_text(report['parser_version']), str(len(records))]) +
            ') ON DUPLICATE KEY UPDATE id=LAST_INSERT_ID(id);',
            'SET @teaching_batch_id=LAST_INSERT_ID();']
    for record in records:
        fields = ['@teaching_batch_id', sql_text(record['sheet']), str(record['source_row']),
                  sql_text(json.dumps(record['raw'], ensure_ascii=False)),
                  sql_text(json.dumps(record['schedule'], ensure_ascii=False)) if record['schedule'] is not None else 'NULL',
                  sql_text(json.dumps(record['issues'], ensure_ascii=False)),
                  sql_text('ERROR' if any(i.startswith('INVALID:') for i in record['issues']) else 'REVIEW')]
        text.append('INSERT INTO teaching_import_row (batch_id,sheet_name,source_row,raw_data,parsed_schedule,issues,status) VALUES (' +
                    ','.join(fields) + ') ON DUPLICATE KEY UPDATE id=id;')
    return '\n'.join(text + ['COMMIT;', ''])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--sql', required=True, type=Path)
    parser.add_argument('--report', required=True, type=Path)
    args = parser.parse_args()
    records, report = analyze(args.workbook)
    for path in [args.sql, args.report]:
        if path.resolve() == args.workbook.resolve():
            parser.error('输出不能覆盖原始 Excel')
        path.parent.mkdir(parents=True, exist_ok=True)
    args.sql.write_text(staging_sql(records, report), encoding='utf-8')
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: report[k] for k in ['row_count', 'course_code_count', 'teacher_name_count',
                     'lab_code_count', 'raw_schedule_segments', 'expanded_schedule_rows', 'issue_counts']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
